"""Dataset loading, resume parsing, and machine-learning analysis helpers."""

from __future__ import annotations

import math
import re
import subprocess
import tempfile
from collections import Counter
from collections.abc import Sequence
from io import BytesIO
from pathlib import Path
from typing import TypedDict
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

import numpy as np
import pandas as pd
from docx import Document
from html.parser import HTMLParser
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader
from pypdf.errors import PdfReadError
import pymupdf
import pytesseract
from striprtf.striprtf import rtf_to_text

try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.linear_model import SGDClassifier
    from sklearn.metrics.pairwise import cosine_similarity
    from sklearn.pipeline import Pipeline
except ImportError:  # pragma: no cover - runtime fallback for environments without the ML stack.
    TfidfVectorizer = None  # type: ignore[assignment]
    SGDClassifier = None  # type: ignore[assignment]
    cosine_similarity = None  # type: ignore[assignment]
    Pipeline = None  # type: ignore[assignment]


SKILL_ALIASES: dict[str, tuple[str, ...]] = {
    "Accounting": ("accounting", "bookkeeping", "quickbooks"),
    "Agile": ("agile", "scrum"),
    "Agriculture": ("agriculture", "crop management", "horticulture"),
    "Adobe Creative Suite": ("adobe photoshop", "photoshop", "illustrator", "indesign"),
    "Aviation": ("aviation", "flight operations", "aviation safety"),
    "AWS": ("aws", "amazon web services"),
    "Azure": ("azure", "microsoft azure"),
    "Banking": ("banking",),
    "Budgeting": ("budgeting", "budget management"),
    "Business Development": ("business development",),
    "CAD": ("cad", "autocad", "revit"),
    "Clinical Care": ("clinical", "patient care", "nursing"),
    "Communication": ("communication", "written communication", "verbal communication"),
    "Compliance": ("compliance", "hipaa", "regulatory compliance"),
    "Construction": ("construction", "blueprint reading", "osha"),
    "Content Writing": ("content writing", "copywriting", "technical writing"),
    "Customer Service": ("customer service", "customer support"),
    "Data Analysis": ("data analysis", "data analytics", "statistical analysis"),
    "Data Visualization": ("data visualization", "tableau", "power bi"),
    "Database Management": ("database management", "mysql", "postgresql"),
    "Digital Marketing": ("digital marketing", "seo", "sem", "social media"),
    "Docker": ("docker",),
    "Education": ("teaching", "curriculum development", "lesson planning"),
    "Electrical Engineering": ("electrical engineering",),
    "Employee Relations": ("employee relations", "human resources"),
    "Engineering": ("engineering", "mechanical engineering", "civil engineering"),
    "Event Planning": ("event planning", "event management"),
    "Excel": ("excel", "microsoft excel", "spreadsheet"),
    "Financial Analysis": ("financial analysis", "financial modeling", "forecasting"),
    "Fitness Training": ("fitness", "personal training", "exercise science"),
    "Food Safety": ("food safety", "food preparation", "menu planning"),
    "Git": ("git", "github"),
    "Graphic Design": ("graphic design", "figma", "ui/ux", "user experience"),
    "Inventory Management": ("inventory management", "inventory control"),
    "Java": ("java",),
    "JavaScript": ("javascript", "react", "node.js", "nodejs"),
    "Jira": ("jira",),
    "Leadership": ("leadership", "team management", "team leadership"),
    "Logistics": ("logistics", "supply chain"),
    "Machine Learning": ("machine learning", "deep learning", "artificial intelligence"),
    "Marketing": ("marketing", "market research", "brand management"),
    "Microsoft Office": ("microsoft office", "word", "powerpoint"),
    "Negotiation": ("negotiation", "negotiating"),
    "Operations": ("operations", "process improvement", "operations management"),
    "Payroll": ("payroll",),
    "Project Management": ("project management", "project planning"),
    "Python": ("python",),
    "Quality Assurance": ("quality assurance", "quality control", "testing"),
    "Recruitment": ("recruitment", "recruiting", "talent acquisition"),
    "Research": ("research", "literature review"),
    "Risk Management": ("risk management", "risk assessment"),
    "SQL": ("sql",),
    "Sales": ("sales", "salesforce", "crm"),
    "Safety": ("safety", "workplace safety"),
    "Tableau": ("tableau",),
    "Tax Preparation": ("tax preparation", "tax compliance"),
    "Training": ("training", "staff development"),
    "Troubleshooting": ("troubleshooting", "technical support"),
    "TypeScript": ("typescript",),
    "UI/UX": ("ui/ux", "user interface", "user experience"),
    "Video Editing": ("video editing", "premiere pro"),
    "Web Development": ("web development", "html", "css", "rest api"),
}

MAX_TRAINING_EXAMPLES_PER_CATEGORY = 500


class HiringOutcomeGroup(TypedDict):
    group: str
    candidates: int
    hired: int
    hiring_rate: float


class HiringDatasetInsights(TypedDict):
    candidate_count: int
    hired_count: int
    not_hired_count: int
    hiring_rate: float
    by_education: list[HiringOutcomeGroup]
    by_company_type: list[HiringOutcomeGroup]
    by_university_tier: list[HiringOutcomeGroup]


class _SimpleTfidfVectorizer:
    """A lightweight TF-IDF fallback used when scikit-learn is unavailable."""

    _STOP_WORDS = {
        "a", "an", "and", "as", "at", "be", "by", "for", "from", "has", "he",
        "in", "is", "it", "its", "of", "on", "that", "the", "to", "was", "were",
        "with", "will", "you", "your", "we", "our", "this", "they", "them",
        "their", "have", "had", "but", "not", "or", "if", "then", "than", "so",
    }

    def __init__(self, *, lowercase=True, strip_accents="unicode", stop_words="english",
                 ngram_range=(1, 2), max_features=50_000, min_df=2, sublinear_tf=True):
        self.lowercase = lowercase
        self.strip_accents = strip_accents
        self.stop_words = stop_words
        self.ngram_range = ngram_range
        self.max_features = max_features
        self.min_df = min_df
        self.sublinear_tf = sublinear_tf
        self.vocabulary_: dict[str, int] = {}
        self.idf_: list[float] = []
        self._document_count = 0

    def _normalize(self, text: str) -> list[str]:
        value = text.casefold() if self.lowercase else text
        terms = re.findall(r"[a-z0-9]+(?:[-/][a-z0-9]+)?", value)
        filtered = []
        for term in terms:
            if self.stop_words == "english" and term in self._STOP_WORDS:
                continue
            filtered.append(term)
        return filtered

    def _get_ngrams(self, terms: list[str]) -> list[str]:
        if self.ngram_range == (1, 1):
            return terms
        ngrams: list[str] = []
        for start in range(len(terms)):
            for end in range(start + 1, min(len(terms), start + self.ngram_range[1]) + 1):
                if end - start < self.ngram_range[0]:
                    continue
                ngrams.append(" ".join(terms[start:end]))
        return ngrams

    def fit(self, documents: Sequence[str]) -> "_SimpleTfidfVectorizer":
        doc_terms: list[list[str]] = []
        document_frequency: Counter[str] = Counter()
        for document in documents:
            terms = self._get_ngrams(self._normalize(str(document)))
            doc_terms.append(terms)
            document_frequency.update(set(terms))
            self._document_count += 1

        selected_terms = [
            term for term, count in document_frequency.items()
            if count >= self.min_df
        ]
        selected_terms = sorted(selected_terms)[: self.max_features]
        self.vocabulary_ = {term: index for index, term in enumerate(selected_terms)}
        self.idf_ = [0.0] * len(self.vocabulary_)

        for term, index in self.vocabulary_.items():
            df = document_frequency.get(term, 0)
            self.idf_[index] = math.log((1 + self._document_count) / (1 + df)) + 1.0

        return self

    def transform(self, documents: Sequence[str]) -> np.ndarray:
        if not self.vocabulary_:
            raise ValueError("The vectorizer has not been fitted yet.")

        rows: list[list[float]] = []
        for document in documents:
            terms = self._get_ngrams(self._normalize(str(document)))
            counts = Counter(terms)
            vector = np.zeros(len(self.vocabulary_), dtype=float)
            for term, count in counts.items():
                index = self.vocabulary_.get(term)
                if index is None:
                    continue
                tf = 1.0 + math.log(count) if self.sublinear_tf else float(count)
                vector[index] = tf * self.idf_[index]
            rows.append(vector)
        return np.asarray(rows, dtype=float)

    def get_feature_names_out(self) -> np.ndarray:
        return np.asarray(list(self.vocabulary_.keys()), dtype=object)


class _SimpleClassifier:
    """Naive linear classifier that provides class probabilities without scikit-learn."""

    def __init__(self, *, class_weight="balanced", random_state=42, n_jobs=1):
        self.class_weight = class_weight
        self.random_state = random_state
        self.n_jobs = n_jobs
        self.classes_: np.ndarray = np.array([], dtype=object)
        self._class_vectors: dict[str, np.ndarray] = {}

    def fit(self, X: np.ndarray, y: Sequence[str]) -> "_SimpleClassifier":
        labels = list(y)
        self.classes_ = np.asarray(sorted(set(labels)), dtype=object)
        weights: dict[str, np.ndarray] = {label: np.zeros(X.shape[1], dtype=float) for label in self.classes_}
        counts: dict[str, int] = {label: 0 for label in self.classes_}

        for row, label in zip(X, labels):
            counts[label] += 1
            weights[label] += np.asarray(row, dtype=float)

        if self.class_weight == "balanced":
            for label in self.classes_:
                class_size = max(counts[label], 1)
                weights[label] = weights[label] / class_size
                if class_size > 1:
                    weights[label] *= (len(self.classes_) / class_size)

        self._class_vectors = weights
        return self

    def predict_proba(self, X: np.ndarray) -> np.ndarray:
        rows = np.asarray(X, dtype=float)
        probabilities = []

        for row in rows:
            scores = []
            for label in self.classes_:
                coefficient = self._class_vectors.get(label, np.zeros(rows.shape[1], dtype=float))
                score = float(np.dot(row, coefficient))
                scores.append(score)

            max_score = max(scores) if scores else 0.0
            if max_score == 0:
                raw = np.full(len(self.classes_), 1.0 / len(self.classes_), dtype=float)
            else:
                raw = np.exp(np.asarray(scores, dtype=float) - max_score)
                raw = raw / raw.sum()
            probabilities.append(raw)

        return np.asarray(probabilities, dtype=float)


class _SimplePipeline:
    def __init__(self):
        self.named_steps = {
            "tfidf": _SimpleTfidfVectorizer(max_features=3_000),
            "classifier": _SimpleClassifier(),
        }

    def fit(self, texts: Sequence[str], labels: Sequence[str]) -> "_SimplePipeline":
        self.named_steps["tfidf"].fit(texts)
        self.named_steps["classifier"].fit(self.named_steps["tfidf"].transform(texts), labels)
        return self

    def predict_proba(self, texts: Sequence[str]) -> np.ndarray:
        transformed = self.named_steps["tfidf"].transform(texts)
        return self.named_steps["classifier"].predict_proba(transformed)


def _to_dense_matrix(matrix: object) -> np.ndarray:
    if hasattr(matrix, "toarray"):
        return np.asarray(matrix.toarray(), dtype=float)
    return np.asarray(matrix, dtype=float)


def _fallback_cosine_similarity(left: np.ndarray, right: np.ndarray) -> float:
    left_array = np.asarray(left, dtype=float).ravel()
    right_array = np.asarray(right, dtype=float).ravel()
    left_norm = np.linalg.norm(left_array)
    right_norm = np.linalg.norm(right_array)
    if left_norm == 0.0 or right_norm == 0.0:
        return 0.0
    return float(np.dot(left_array, right_array) / (left_norm * right_norm))


def summarize_hiring_outcomes(data: pd.DataFrame) -> HiringDatasetInsights:
    """Summarize binary hiring outcomes without exposing candidate-level records."""
    group_columns = ("education_level", "company_type", "university_tier")
    required_columns = {"hired", *group_columns}
    missing_columns = required_columns.difference(data.columns)
    if missing_columns:
        raise ValueError(
            "Hiring dataset is missing required columns: " + ", ".join(sorted(missing_columns))
        )
    if data.empty:
        raise ValueError("Hiring dataset contains no candidate records.")

    summary_data = data.loc[:, ["hired", *group_columns]].copy()
    summary_data["hired"] = pd.to_numeric(summary_data["hired"], errors="coerce")
    if summary_data.isna().any().any():
        raise ValueError("Hiring dataset contains missing or invalid outcome/group values.")
    if not summary_data["hired"].isin([0, 1]).all():
        raise ValueError("Hiring dataset's hired column must contain only 0 and 1.")

    hired_count = int(summary_data["hired"].sum())

    def summarize_group(column: str) -> list[HiringOutcomeGroup]:
        grouped = (
            summary_data.groupby(column, sort=True, observed=True)["hired"]
            .agg(candidates="size", hired="sum")
            .reset_index()
        )
        return [
            {
                "group": str(row[column]),
                "candidates": int(row["candidates"]),
                "hired": int(row["hired"]),
                "hiring_rate": float(row["hired"] / row["candidates"]),
            }
            for _, row in grouped.iterrows()
        ]

    return {
        "candidate_count": len(summary_data),
        "hired_count": hired_count,
        "not_hired_count": len(summary_data) - hired_count,
        "hiring_rate": float(hired_count / len(summary_data)),
        "by_education": summarize_group("education_level"),
        "by_company_type": summarize_group("company_type"),
        "by_university_tier": summarize_group("university_tier"),
    }


def _normalize_csv_paths(csv_path: str | Path | Sequence[str | Path] | None) -> list[str | Path]:
    """Return a flat list of CSV sources from a path or a collection of paths."""
    if csv_path is None:
        return []
    if isinstance(csv_path, (str, Path)):
        return [csv_path]
    return list(csv_path)


def load_training_data(
    csv_path: str | Path | Sequence[str | Path] | None = None,
    pdf_root: str | Path | None = None,
    csv_bytes: bytes | None = None,
) -> tuple[pd.DataFrame, dict[str, int]]:
    """Load labeled CSV resumes and, when available, resumes in category PDF folders."""
    csv_frames: list[pd.DataFrame] = []
    csv_sources_loaded = 0
    csv_sources_skipped = 0
    csv_duplicates_removed = 0

    if csv_bytes is not None:
        csv_sources = [BytesIO(csv_bytes)]
    else:
        csv_sources = []
        for csv_entry in _normalize_csv_paths(csv_path):
            csv_file = Path(csv_entry).expanduser()
            if not csv_file.is_file():
                csv_sources_skipped += 1
                continue
            csv_sources.append(csv_file)

    if not csv_sources and csv_bytes is None:
        raise ValueError("Provide a CSV path or upload a CSV file.")

    for csv_source in csv_sources:
        try:
            csv_header = pd.read_csv(csv_source, nrows=0, encoding_errors="replace")
        except (OSError, ValueError, pd.errors.ParserError):
            csv_sources_skipped += 1
            continue

        normalized_columns = {
            str(column).strip().casefold().replace(" ", "_"): str(column)
            for column in csv_header.columns
        }
        text_column = next(
            (
                normalized_columns[alias]
                for alias in ("resume_str", "resume", "feature", "text", "resume_text")
                if alias in normalized_columns
            ),
            None,
        )
        label_column = next(
            (
                normalized_columns[alias]
                for alias in ("category", "label", "job_category")
                if alias in normalized_columns
            ),
            None,
        )
        if text_column is None or label_column is None:
            csv_sources_skipped += 1
            continue

        if isinstance(csv_source, BytesIO):
            csv_source.seek(0)
        try:
            csv_data = pd.read_csv(
                csv_source,
                usecols=[text_column, label_column],
                encoding_errors="replace",
            )
        except (OSError, ValueError, pd.errors.ParserError):
            csv_sources_skipped += 1
            continue

        csv_subset = csv_data.rename(columns={text_column: "text", label_column: "category"})
        csv_subset["text"] = csv_subset["text"].fillna("").astype(str).str.strip()
        csv_subset["category"] = csv_subset["category"].fillna("").astype(str).str.strip()
        csv_subset = csv_subset.loc[
            csv_subset["text"].ne("") & csv_subset["category"].ne(""),
            ["text", "category"],
        ]
        if label_column.casefold() == "label":
            csv_subset["category"] = csv_subset["category"].str.split(",")
            csv_subset = csv_subset.explode("category")
            csv_subset["category"] = csv_subset["category"].str.strip()
            csv_subset = csv_subset.loc[csv_subset["category"].ne("")]
        if not csv_subset.empty:
            csv_frames.append(csv_subset)
            csv_sources_loaded += 1
        else:
            csv_sources_skipped += 1

    if not csv_frames:
        raise ValueError("No usable labeled resumes were found in the selected CSV files.")

    csv_data = pd.concat(csv_frames, ignore_index=True)
    csv_count_before_deduplication = len(csv_data)
    csv_data = csv_data.drop_duplicates(subset=["text", "category"], ignore_index=True)
    csv_duplicates_removed = csv_count_before_deduplication - len(csv_data)

    pdf_documents: list[dict[str, str]] = []
    pdf_failures = 0
    if pdf_root:
        root = Path(pdf_root).expanduser()
        if root.is_dir():
            for pdf_path in root.rglob("*.pdf"):
                category = pdf_path.parent.name.strip()
                try:
                    reader = PdfReader(pdf_path)
                    text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
                except (OSError, ValueError, PdfReadError):
                    pdf_failures += 1
                    continue
                if category and text:
                    pdf_documents.append({"text": text, "category": category})

    pdf_data = pd.DataFrame(pdf_documents, columns=["text", "category"])
    combined = pd.concat([csv_data, pdf_data], ignore_index=True)
    stats = {
        "csv_resumes": len(csv_data),
        "pdf_resumes": len(pdf_data),
        "pdf_failures": pdf_failures,
        "categories": combined["category"].nunique(),
        "csv_sources_loaded": csv_sources_loaded,
        "csv_sources_skipped": csv_sources_skipped,
        "csv_duplicates_removed": csv_duplicates_removed,
    }
    return combined, stats


def _select_training_examples(training_data: pd.DataFrame) -> pd.DataFrame:
    """Keep the training set balanced and bounded while retaining every category."""
    category_groups = training_data.groupby("category", sort=True, observed=True)
    selected = [
        group.sample(
            n=MAX_TRAINING_EXAMPLES_PER_CATEGORY,
            random_state=42,
        )
        if len(group) > MAX_TRAINING_EXAMPLES_PER_CATEGORY
        else group
        for _, group in category_groups
    ]
    return pd.concat(selected, ignore_index=True)


def train_classifier(training_data: pd.DataFrame):
    """Fit a bounded TF-IDF classifier on the labeled resumes."""
    if training_data.empty:
        raise ValueError("No usable labeled resumes were found in the selected datasets.")

    training_data = _select_training_examples(training_data)

    if TfidfVectorizer is None or SGDClassifier is None or Pipeline is None:
        model = _SimplePipeline()
        model.fit(training_data["text"].tolist(), training_data["category"].tolist())
        model.training_example_count_ = len(training_data)
        return model

    model = Pipeline(
        [
            (
                "tfidf",
                TfidfVectorizer(
                    lowercase=True,
                    strip_accents="unicode",
                    stop_words="english",
                    ngram_range=(1, 2),
                    max_features=50_000,
                    min_df=2,
                    sublinear_tf=True,
                ),
            ),
            (
                "classifier",
                SGDClassifier(
                    loss="log_loss",
                    max_iter=30,
                    tol=1e-3,
                    class_weight="balanced",
                    random_state=42,
                    n_jobs=1,
                ),
            ),
        ]
    )
    model.fit(training_data["text"].tolist(), training_data["category"].tolist())
    model.training_example_count_ = len(training_data)
    return model


class _VisibleHTMLText(HTMLParser):
    """Collect visible text from HTML while ignoring script and style contents."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_tags: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style"}:
            self._ignored_tags.append(tag.lower())

    def handle_endtag(self, tag: str) -> None:
        if self._ignored_tags and tag.lower() == self._ignored_tags[-1]:
            self._ignored_tags.pop()

    def handle_data(self, data: str) -> None:
        if not self._ignored_tags:
            self.parts.append(data)


def _extract_ocr_image(image: Image.Image, filename: str) -> str:
    try:
        return pytesseract.image_to_string(image)
    except pytesseract.TesseractNotFoundError as error:
        raise ValueError(
            f"OCR is not available for {filename}. Try uploading a text-based PDF or DOCX instead."
        ) from error


def _extract_doc_text(content: bytes) -> str:
    try:
        with tempfile.NamedTemporaryFile(suffix=".doc") as temporary_file:
            temporary_file.write(content)
            temporary_file.flush()
            result = subprocess.run(
                ["antiword", temporary_file.name],
                check=True,
                capture_output=True,
                timeout=30,
            )
    except FileNotFoundError as error:
        raise ValueError(
            "Legacy .doc files need a converter. Save the resume as .docx or PDF and upload it again."
        ) from error
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as error:
        raise ValueError("The .doc file could not be converted. Try saving it as .docx or PDF.") from error
    return result.stdout.decode("utf-8", errors="replace")


def extract_resume_text(filename: str, content: bytes) -> str:
    """Extract text or OCR from common document, image, and text upload formats."""
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".text", ".md", ".markdown", ".csv", ".json"}:
        text = content.decode("utf-8", errors="replace")
    elif suffix == ".pdf":
        try:
            reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except PdfReadError as error:
            raise ValueError("The PDF could not be read. Try exporting it again.") from error
        if not text.strip():
            try:
                with pymupdf.open(stream=content, filetype="pdf") as document:
                    if len(document) > 30:
                        raise ValueError(
                            "This scanned PDF has more than 30 pages. Run OCR on it or split it into smaller files."
                        )
                    extracted_pages = []
                    for page in document:
                        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
                        image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                        extracted_pages.append(_extract_ocr_image(image, filename))
                    text = "\n".join(extracted_pages)
            except (pymupdf.FileDataError, pymupdf.EmptyFileError) as error:
                raise ValueError("The PDF could not be read. Try exporting it again.") from error
    elif suffix == ".docx":
        try:
            document = Document(BytesIO(content))
        except BadZipFile as error:
            raise ValueError("The DOCX file could not be read. Try exporting it again.") from error
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        text = "\n".join(parts)
    elif suffix == ".doc":
        text = _extract_doc_text(content)
    elif suffix == ".rtf":
        text = rtf_to_text(content.decode("utf-8", errors="replace"))
    elif suffix == ".odt":
        try:
            with ZipFile(BytesIO(content)) as document:
                root = ElementTree.fromstring(document.read("content.xml"))
        except (BadZipFile, KeyError, ElementTree.ParseError) as error:
            raise ValueError("The ODT file could not be read. Try exporting it again.") from error
        text = "\n".join(
            "".join(element.itertext())
            for element in root.iter()
            if element.tag.rsplit("}", 1)[-1] in {"p", "h"}
        )
    elif suffix in {".html", ".htm"}:
        parser = _VisibleHTMLText()
        parser.feed(content.decode("utf-8", errors="replace"))
        text = " ".join(parser.parts)
    elif suffix in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}:
        try:
            with Image.open(BytesIO(content)) as image:
                text = _extract_ocr_image(image, filename)
        except UnidentifiedImageError as error:
            raise ValueError("The image could not be opened. Try uploading a PNG or JPEG image.") from error
    elif suffix in {".xlsx", ".xls"}:
        try:
            workbook = pd.read_excel(BytesIO(content), header=None, dtype=str)
        except (ImportError, ValueError, OSError) as error:
            raise ValueError("The spreadsheet could not be read. Try exporting it as .xlsx or PDF.") from error
        text = "\n".join(
            " ".join(str(value) for value in row if pd.notna(value))
            for row in workbook.itertuples(index=False, name=None)
        )
    else:
        raise ValueError(
            f"The {suffix or 'unknown'} file type is not supported for text extraction. "
            "Try PDF, DOC/DOCX, ODT, RTF, TXT, HTML, CSV, XLS/XLSX, or an image file."
        )

    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        raise ValueError("No readable text was found. Check that the file contains text and is not blank.")
    return text


def extract_skills(text: str) -> set[str]:
    """Return skill labels found using case-insensitive whole-phrase matching."""
    normalized = text.casefold()
    found: set[str] = set()
    for skill, aliases in SKILL_ALIASES.items():
        for alias in aliases:
            pattern = rf"(?<!\w){re.escape(alias.casefold())}(?!\w)"
            if re.search(pattern, normalized):
                found.add(skill)
                break
    return found


def analyze_resume(
    model: Pipeline,
    resume_text: str,
    job_description: str = "",
) -> dict[str, object]:
    """Predict suitable dataset categories and optionally score a job-description match."""
    classifier = model.named_steps["classifier"]
    vectorizer = model.named_steps["tfidf"]
    probabilities = model.predict_proba([resume_text])[0]
    ranked_indices = probabilities.argsort()[::-1][:3]
    recommendations = [
        {
            "category": str(classifier.classes_[index]),
            "confidence": float(probabilities[index]),
        }
        for index in ranked_indices
    ]

    resume_skills = extract_skills(resume_text)
    result: dict[str, object] = {
        "recommendations": recommendations,
        "resume_skills": sorted(resume_skills),
    }

    job_description = job_description.strip()
    if not job_description:
        return result

    job_skills = extract_skills(job_description)
    matched_skills = sorted(resume_skills & job_skills)
    missing_skills = sorted(job_skills - resume_skills)

    resume_vector = vectorizer.transform([resume_text])
    job_vector = vectorizer.transform([job_description])
    if cosine_similarity is not None:
        similarity = float(cosine_similarity(resume_vector, job_vector)[0, 0])
    else:
        similarity = _fallback_cosine_similarity(resume_vector[0], job_vector[0])

    if job_skills:
        coverage = len(matched_skills) / len(job_skills)
        alignment = 100 * (0.75 * coverage + 0.25 * similarity)
    else:
        coverage = None
        alignment = 100 * similarity

    feature_names = vectorizer.get_feature_names_out()
    resume_dense = _to_dense_matrix(resume_vector)
    job_dense = _to_dense_matrix(job_vector)
    job_weights = job_dense[0]
    resume_terms = set(feature_names[resume_dense[0] > 0])
    ranked_terms = job_weights.argsort()[::-1]
    missing_keywords: list[str] = []
    for index in ranked_terms:
        term = str(feature_names[index])
        if job_weights[index] <= 0 or term in resume_terms:
            continue
        if any(term in keyword or keyword in term for keyword in missing_keywords):
            continue
        missing_keywords.append(term)
        if len(missing_keywords) == 8:
            break

    result.update(
        {
            "job_skills": sorted(job_skills),
            "matched_skills": matched_skills,
            "missing_skills": missing_skills,
            "skill_coverage": coverage,
            "text_similarity": similarity,
            "alignment": min(100.0, max(0.0, alignment)),
            "missing_keywords": missing_keywords,
        }
    )
    return result
