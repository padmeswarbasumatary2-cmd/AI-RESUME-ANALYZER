from io import BytesIO
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
import pymupdf
from docx import Document
from openpyxl import Workbook

from analyzer import (
    analyze_resume,
    extract_resume_text,
    extract_skills,
    load_training_data,
    summarize_hiring_outcomes,
    train_classifier,
)


class AnalyzerTests(unittest.TestCase):
    def setUp(self):
        self.training_data = pd.DataFrame(
            [
                {
                    "text": "Python SQL data analysis machine learning dashboards",
                    "category": "INFORMATION-TECHNOLOGY",
                },
                {
                    "text": "Python SQL software engineering data analysis",
                    "category": "INFORMATION-TECHNOLOGY",
                },
                {
                    "text": "Python data analytics and software development SQL",
                    "category": "INFORMATION-TECHNOLOGY",
                },
                {
                    "text": "patient care clinical nursing healthcare records",
                    "category": "HEALTHCARE",
                },
                {
                    "text": "healthcare nursing patient care medical records",
                    "category": "HEALTHCARE",
                },
                {
                    "text": "clinical healthcare patient support and nursing",
                    "category": "HEALTHCARE",
                },
            ]
        )
        self.model = train_classifier(self.training_data)

    def test_skill_matching_uses_whole_terms(self):
        skills = extract_skills("Python and SQL skills; experience with Pythonic tooling.")

        self.assertIn("Python", skills)
        self.assertIn("SQL", skills)

    def test_analyzer_ranks_related_category_and_job_skills(self):
        result = analyze_resume(
            self.model,
            "Python SQL data analysis software engineering",
            "We need Python, SQL, and project management experience.",
        )

        self.assertEqual(result["recommendations"][0]["category"], "INFORMATION-TECHNOLOGY")
        self.assertIn("Python", result["matched_skills"])
        self.assertIn("SQL", result["matched_skills"])
        self.assertIn("Project Management", result["missing_skills"])
        self.assertGreaterEqual(result["alignment"], 0)
        self.assertLessEqual(result["alignment"], 100)

    def test_classifier_caps_each_category_without_dropping_categories(self):
        training_data = pd.concat(
            [
                pd.DataFrame(
                    {
                        "text": [f"Python SQL software project {index}" for index in range(510)],
                        "category": ["TECH"] * 510,
                    }
                ),
                pd.DataFrame(
                    {
                        "text": ["patient care clinical nursing"] * 3,
                        "category": ["HEALTHCARE"] * 3,
                    }
                ),
            ],
            ignore_index=True,
        )

        model = train_classifier(training_data)

        self.assertEqual(model.training_example_count_, 503)
        self.assertEqual(set(model.named_steps["classifier"].classes_), {"HEALTHCARE", "TECH"})

    def test_txt_resume_extraction_and_empty_document_error(self):
        text = extract_resume_text("resume.txt", b"  Python   developer ")

        self.assertEqual(text, "Python developer")
        with self.assertRaisesRegex(ValueError, "No readable text"):
            extract_resume_text("empty.txt", b" \n ")

    def test_text_pdf_resume_extraction(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Python developer with SQL experience")
        pdf_content = document.tobytes()
        document.close()

        text = extract_resume_text("resume.pdf", pdf_content)

        self.assertIn("Python developer", text)
        self.assertIn("SQL experience", text)

    def test_rtf_resume_extraction(self):
        text = extract_resume_text("resume.rtf", br"{\rtf1\ansi Python developer}")

        self.assertIn("Python developer", text)

    def test_html_extraction_ignores_script_and_styles(self):
        text = extract_resume_text(
            "resume.html",
            b"<h1>Python Engineer</h1><script>ignore this</script><p>SQL skills</p>",
        )

        self.assertIn("Python Engineer", text)
        self.assertIn("SQL skills", text)
        self.assertNotIn("ignore this", text)

    def test_odt_resume_extraction(self):
        odt = BytesIO()
        with ZipFile(odt, "w") as archive:
            archive.writestr(
                "content.xml",
                '<document xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
                "<text:p>Python Developer</text:p><text:p>SQL analyst</text:p></document>",
            )

        text = extract_resume_text("resume.odt", odt.getvalue())

        self.assertIn("Python Developer", text)
        self.assertIn("SQL analyst", text)

    def test_xlsx_resume_extraction(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["Python developer", "SQL analyst"])
        content = BytesIO()
        workbook.save(content)

        text = extract_resume_text("resume.xlsx", content.getvalue())

        self.assertIn("Python developer", text)
        self.assertIn("SQL analyst", text)

    def test_unsupported_format_has_actionable_error(self):
        with self.assertRaisesRegex(ValueError, "file type is not supported"):
            extract_resume_text("resume.pages", b"binary document")

    def test_invalid_docx_is_reported_as_a_readable_error(self):
        with self.assertRaisesRegex(ValueError, "DOCX file could not be read"):
            extract_resume_text("resume.docx", b"not a document")

    def test_docx_resume_extraction(self):
        document = Document()
        document.add_paragraph("Python developer with SQL experience")
        content = BytesIO()
        document.save(content)

        text = extract_resume_text("resume.docx", content.getvalue())

        self.assertIn("Python developer", text)
        self.assertIn("SQL experience", text)

    def test_csv_loader_uses_labeled_resume_text(self):
        with tempfile.TemporaryDirectory() as directory:
            csv_path = Path(directory) / "resumes.csv"
            pd.DataFrame(
                {
                    "Resume_str": ["Python developer", "Nurse"],
                    "Category": ["INFORMATION-TECHNOLOGY", "HEALTHCARE"],
                }
            ).to_csv(csv_path, index=False)

            data, stats = load_training_data(csv_path)

        self.assertEqual(len(data), 2)
        self.assertEqual(stats["csv_resumes"], 2)
        self.assertEqual(stats["pdf_resumes"], 0)
        self.assertEqual(stats["categories"], 2)

    def test_csv_loader_accepts_in_memory_upload(self):
        csv_bytes = (
            b"Resume_str,Category\n"
            b'"Python data analysis",INFORMATION-TECHNOLOGY\n'
            b'"Clinical patient care",HEALTHCARE\n'
        )

        data, stats = load_training_data(csv_bytes=csv_bytes)

        self.assertEqual(len(data), 2)
        self.assertEqual(stats["csv_resumes"], 2)
        self.assertEqual(stats["categories"], 2)

    def test_csv_loader_ignores_non_resume_csvs_and_keeps_valid_ones(self):
        with tempfile.TemporaryDirectory() as directory:
            valid_path = Path(directory) / "valid.csv"
            invalid_path = Path(directory) / "candidate_metrics.csv"
            pd.DataFrame(
                {
                    "Resume_str": ["Python developer", "Nurse"],
                    "Category": ["INFORMATION-TECHNOLOGY", "HEALTHCARE"],
                }
            ).to_csv(valid_path, index=False)
            pd.DataFrame({"candidate_id": [1], "hired": [1], "skills_score": [20]}).to_csv(
                invalid_path,
                index=False,
            )

            data, stats = load_training_data([valid_path, invalid_path])

        self.assertEqual(len(data), 2)
        self.assertEqual(stats["csv_resumes"], 2)
        self.assertEqual(stats["csv_sources_loaded"], 1)
        self.assertGreaterEqual(stats["csv_sources_skipped"], 1)

    def test_csv_loader_supports_resume_aliases_gzip_and_multi_category_labels(self):
        with tempfile.TemporaryDirectory() as directory:
            directory_path = Path(directory)
            resume_path = directory_path / "resume.csv"
            feature_path = directory_path / "feature.csv"
            labels_path = directory_path / "labels.csv.gz"
            pd.DataFrame(
                {
                    "Category": ["DATA", "TECH"],
                    "Resume": ["Shared resume", "Data analysis resume"],
                }
            ).to_csv(resume_path, index=False)
            pd.DataFrame(
                {
                    "ID": [1, 2],
                    "Category": ["DATA", "FINANCE"],
                    "Feature": ["Shared resume", "Bookkeeping resume"],
                }
            ).to_csv(feature_path, index=False)
            pd.DataFrame(
                {
                    "Resume": ["Python software engineering resume"],
                    "Label": ["Python_Developer,Software_Developer"],
                }
            ).to_csv(labels_path, index=False, compression="gzip")

            data, stats = load_training_data([resume_path, feature_path, labels_path])

        self.assertEqual(len(data), 5)
        self.assertEqual(stats["csv_resumes"], 5)
        self.assertEqual(stats["csv_sources_loaded"], 3)
        self.assertEqual(stats["csv_duplicates_removed"], 1)
        self.assertIn("Python_Developer", set(data["category"]))
        self.assertIn("Software_Developer", set(data["category"]))

    def test_hiring_dataset_summary_returns_only_aggregate_outcomes(self):
        summary = summarize_hiring_outcomes(
            pd.DataFrame(
                {
                    "candidate_id": [101, 102, 103],
                    "age": [22, 27, 31],
                    "education_level": ["Bachelors", "Masters", "Bachelors"],
                    "company_type": ["Startup", "MNC", "Startup"],
                    "university_tier": ["Tier 1", "Tier 2", "Tier 1"],
                    "hired": [1, 0, 1],
                }
            )
        )

        self.assertEqual(summary["candidate_count"], 3)
        self.assertEqual(summary["hired_count"], 2)
        self.assertAlmostEqual(summary["hiring_rate"], 2 / 3)
        self.assertEqual(summary["by_education"][0]["group"], "Bachelors")
        self.assertEqual(summary["by_education"][0]["candidates"], 2)
        self.assertNotIn("age", summary)
        self.assertNotIn("candidate_id", summary)

    def test_hiring_dataset_summary_rejects_non_binary_outcomes(self):
        with self.assertRaisesRegex(ValueError, "only 0 and 1"):
            summarize_hiring_outcomes(
                pd.DataFrame(
                    {
                        "hired": [2],
                        "education_level": ["Bachelors"],
                        "company_type": ["Startup"],
                        "university_tier": ["Tier 1"],
                    }
                )
            )


if __name__ == "__main__":
    unittest.main()
